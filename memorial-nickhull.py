import streamlit as st
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# ==========================================
# CONFIGURAÇÕES DA PÁGINA
# ==========================================
st.set_page_config(page_title="Memorial Descritivo PRO", page_icon="📐", layout="wide")

# ==========================================
# CABEÇALHO OFICIAL COM MARKETING ESTRATÉGICO
# ==========================================
st.markdown("""
    <div style='background-color: #e5e8e8; padding: 25px 20px; border-radius: 8px; text-align: center; border: 1px solid #cccccc; box-shadow: 0 4px 6px rgba(0,0,0,0.1); margin-bottom: 20px;'>
        <h1 style='color: #1f77b4; margin: 0; font-size: 26px; font-weight: bold; text-transform: uppercase; letter-spacing: 1px;'>
            NICK HULL EMERSON ENGINEERING
        </h1>
        <p style='color: #333333; font-style: italic; font-size: 16px; margin: 8px 0 15px 0;'>
            Soluções de Baixo Atrito para Ativos Complexos — Unindo Rigor Técnico e Inteligência Exponencial
        </p>
        <p style='color: #555555; font-size: 14px; margin: 0;'>
            Eng. Civil e Eng. de Seg. do Trabalho <b>Emerson Alves dos Santos</b> | Pós-graduado em Perícias da Engenharia
        </p>
        <p style='color: #555555; font-size: 14px; margin: 2px 0 15px 0;'>
            CREA: 506.999.155-2 | Email: nickhull.eng@gmail.com | Tel: 55 11 99851-1552
        </p>
        <div style='background-color: #ffffff; border-left: 4px solid #1f77b4; padding: 10px; text-align: left; border-radius: 4px;'>
            <p style='margin: 0; font-size: 14px; color: #222;'>
                🚀 <b>Mentoria e Consultoria para Profissionais:</b> Prestamos suporte técnico e estratégico para engenheiros, arquitetos e topógrafos na superação de exigências complexas em Cartórios e Prefeituras. <b>Destrave seus processos e garanta segurança jurídica aos seus clientes.</b>
            </p>
        </div>
    </div>
""", unsafe_allow_html=True)

st.title("Gerador de Memorial Descritivo Oficial")

# ==========================================
# CONTROLE DE ESTADO (VÉRTICES)
# ==========================================
if 'num_vertices' not in st.session_state:
    st.session_state.num_vertices = 4

def adicionar_vertice():
    st.session_state.num_vertices += 1

def remover_vertice():
    if st.session_state.num_vertices > 4:
        st.session_state.num_vertices -= 1

# ==========================================
# FORMULÁRIO DE ENTRADA
# ==========================================
with st.expander("1. DADOS DE LOCALIZAÇÃO DO LOTE", expanded=False):
    col1, col2, col3 = st.columns(3)
    lote_end = col1.text_input("Lote situado à", value="")
    lote_num = col2.text_input("Número", value="")
    bairro = col3.text_input("Bairro", value="")
    cidade = col1.text_input("Cidade", value="")
    estado = col2.text_input("Estado (Sigla)", value="", max_chars=2)
    
    st.write("No quarteirão completado por (Logradouros):")
    q_rua1 = col1.text_input("Rua/Av. 1", value="")
    q_rua2 = col2.text_input("Rua/Av. 2", value="")
    q_rua3 = col3.text_input("Rua/Av. 3", value="")
    q_rua4 = col1.text_input("Rua/Av. 4", value="")

with st.expander("2. PONTO DE AMARRAÇÃO E INÍCIO (PONTO 01)", expanded=False):
    col1, col2 = st.columns(2)
    p1_e = col1.text_input("Coordenada E (Ponto 01)", value="")
    p1_n = col2.text_input("Coordenada N (Ponto 01)", value="")
    p1_alinhamento = col1.text_input("No alinhamento da (Rua/Av)", value="")
    p1_dist_amarra = col2.text_input("Distância do ponto de amarração (m)", value="")
    amarra_e = col1.text_input("Coord. E (Amarração)", value="")
    amarra_n = col2.text_input("Coord. N (Amarração)", value="")
    intersecao_rua1 = col1.text_input("Interseção Rua/Av A", value="")
    intersecao_rua2 = col2.text_input("Interseção Rua/Av B", value="")

st.header("3. DESCRITIVO DOS TRECHOS E CONFRONTANTES")
trechos_dados = []

for i in range(1, st.session_state.num_vertices + 1):
    ponto_atual = i
    ponto_proximo = i + 1 if i < st.session_state.num_vertices else 1
    
    st.subheader(f"Trecho: Ponto {ponto_atual:02d} ao {ponto_proximo:02d}")
    
    colA, colB = st.columns(2)
    
    deflexao, ang_interno = "", ""
    if ponto_atual > 1:
        deflexao = colA.selectbox(f"Deflexão (Ponto {ponto_atual:02d})", ["esquerda", "direita", "continua no alinhamento"], key=f"def_{i}")
        ang_interno = colB.text_input(f"Formando Ângulo Interno de", value="", key=f"ang_{i}")
    else:
        st.caption("Ponto 01 (Início do caminhamento)")
        
    azimute = colA.text_input(f"Azimute", value="", key=f"az_{i}")
    distancia = colB.number_input(f"Distância Total do Trecho (m)", min_value=0.0, value=0.0, format="%.2f", key=f"dist_{i}")
    
    if ponto_atual < st.session_state.num_vertices:
        coord_e = colA.text_input(f"Coordenada E (Ponto {ponto_proximo:02d})", value="", key=f"ce_{i}")
        coord_n = colB.text_input(f"Coordenada N (Ponto {ponto_proximo:02d})", value="", key=f"cn_{i}")
    else:
        coord_e, coord_n = p1_e, p1_n

    st.markdown("**Confrontantes deste trecho:**")
    tipo_trecho = st.radio("Quantidade de confrontantes:", ["Único Confrontante", "Múltiplos Confrontantes"], horizontal=True, key=f"tipo_t_{i}")
    
    textos_confrontantes = []
    
    if tipo_trecho == "Único Confrontante":
        colC1, colC2 = st.columns(2)
        tipo_conf = colC1.selectbox("Tipo", ["Particular", "Logradouro Público", "Mesmo alinhamento (Logradouro)"], key=f"tconf_{i}")
        
        if tipo_conf == "Particular":
            conf_num = colC2.text_input("N° do Imóvel", value="", key=f"cnum_{i}")
            conf_rua = colC1.text_input("Rua/Av do Imóvel", value="", key=f"crua_{i}")
            conf_mat = colC2.text_input("Matrícula", value="", key=f"cmat_{i}")
            textos_confrontantes.append(f"confrontando neste trecho com o imóvel de n° {conf_num}, da {conf_rua} (objeto da matrícula n° {conf_mat} deste registro de imóveis)")
        elif tipo_conf == "Logradouro Público":
            conf_rua_pub = colC2.text_input("Nome da Rua/Av", value="", key=f"cruapub_{i}")
            textos_confrontantes.append(f"confrontando neste trecho com a {conf_rua_pub}")
        else:
            conf_rua_pub = colC2.text_input("Nome da Rua/Av", value="", key=f"cruamesmo_{i}")
            textos_confrontantes.append(f"confrontando no mesmo alinhamento da {conf_rua_pub}")
            
    else:
        num_confrontantes = st.number_input("Quantos confrontantes dividem este trecho?", min_value=2, max_value=10, value=2, key=f"num_div_{i}")
        soma_distancias = 0.0
        
        for j in range(int(num_confrontantes)):
            st.markdown(f"**Confrontante {j+1}**")
            colX, colY, colZ1, colZ2 = st.columns([1, 1, 1.5, 1.5])
            
            dist_frag = colX.number_input("Distância (m)", min_value=0.0, value=0.0, format="%.2f", key=f"dfrag_{i}_{j}")
            soma_distancias += dist_frag 
            
            tipo_conf_frag = colY.selectbox("Tipo", ["Particular", "Logradouro Público", "Mesmo alinhamento (Logradouro)"], key=f"tconffrag_{i}_{j}")
            
            if tipo_conf_frag == "Particular":
                det_imovel = colZ1.text_input("Imóvel (N° e Rua)", value="", placeholder="Ex: imóvel de n° 10 da Rua X", key=f"detimovel_{i}_{j}")
                det_mat = colZ2.text_input("Matrícula", value="", key=f"detmat_{i}_{j}")
                textos_confrontantes.append(f"por {dist_frag:.2f}m com o {det_imovel} (objeto da matrícula n° {det_mat} deste registro de imóveis)")
            elif tipo_conf_frag == "Logradouro Público":
                det_pub = colZ1.text_input("Nome do Logradouro", value="", key=f"detpub_{i}_{j}")
                textos_confrontantes.append(f"por {dist_frag:.2f}m com a {det_pub}")
            else:
                det_pub = colZ1.text_input("Nome do Logradouro", value="", key=f"detpubmesmo_{i}_{j}")
                textos_confrontantes.append(f"por {dist_frag:.2f}m no mesmo alinhamento da {det_pub}")
        
        if abs(soma_distancias - distancia) > 0.01 and distancia > 0:
            st.error(f"⚠️ **ALERTA DE RIGOR FORENSE:** A soma das distâncias dos confrontantes ({soma_distancias:.2f}m) NÃO BATE com a distância total do trecho ({distancia:.2f}m). Verifique as medidas.")
        elif abs(soma_distancias - distancia) <= 0.01 and distancia > 0:
            st.success(f"✅ Fechamento matemático do trecho validado ({soma_distancias:.2f}m).")

    if tipo_trecho == "Único Confrontante":
        txt_final_confrontante = textos_confrontantes[0] if textos_confrontantes else ""
    else:
        txt_final_confrontante = "confrontando neste trecho da seguinte forma: " + ", ".join(textos_confrontantes[:-1]) + (" e " + textos_confrontantes[-1] if len(textos_confrontantes) > 1 else "".join(textos_confrontantes))

    trechos_dados.append({
        "ponto_atual": ponto_atual, "ponto_proximo": ponto_proximo,
        "deflexao": deflexao, "ang_interno": ang_interno, "azimute": azimute,
        "distancia": distancia, "coord_e": coord_e, "coord_n": coord_n,
        "txt_confrontante": txt_final_confrontante
    })
    st.divider()

col_btn1, col_btn2, _ = st.columns([1, 1, 4])
col_btn1.button("➕ Adicionar Vértice", on_click=adicionar_vertice)
col_btn2.button("➖ Remover Vértice", on_click=remover_vertice)

with st.expander("4. DADOS DE FECHAMENTO E ASSINATURA", expanded=False):
    col1, col2 = st.columns(2)
    ang_fechamento = col1.text_input("Ângulo Interno de Fechamento (Ponto 01)", value="")
    perimetro = col1.text_input("Perímetro Total (m)", value="")
    area_total = col2.text_input("Área Total (m²)", value="")
    art = col1.text_input("Número da ART / TRT", value="")
    proprietario = col2.text_input("Nome do Proprietário", value="")
    rg = col1.text_input("RG do Proprietário", value="")
    cpf = col2.text_input("CPF do Proprietário", value="")

# ==========================================
# MOTOR DE GERAÇÃO DO TEXTO (WORD) - ABNT/CARTORIAL
# ==========================================
def gerar_documento():
    doc = Document()
    
    # Configuração das Margens ABNT/Rigor Forense
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    titulo = doc.add_paragraph()
    titulo_run = titulo.add_run("MEMORIAL DESCRITIVO")
    titulo_run.bold = True
    titulo_run.underline = True
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    doc.add_paragraph()
    
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    texto_loc = f"Lote situado à {lote_end}, número {lote_num} – {bairro} – {cidade} – {estado}. No quarteirão completado pela {q_rua1}, {q_rua2}, {q_rua3} e {q_rua4}."
    p1.add_run(texto_loc)
    
    doc.add_paragraph()
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    texto_desc = f"Descrição: Tem início no ponto 01 e possui coordenadas Relativas, Sistema UTM Datum SIRGAS 2000, E= {p1_e}m e N= {p1_n}m, referentes ao Meridiano Central 45 WGr, no alinhamento da {p1_alinhamento}, distante {p1_dist_amarra}m do ponto de amarração de coordenadas E= {amarra_e}m e N= {amarra_n}m, da interseção formada pela {intersecao_rua1} e {intersecao_rua2}; "
    
    for t in trechos_dados:
        if t["deflexao"] == "continua no alinhamento":
            frase_deflexao = f"deste ponto, continua no alinhamento, formando um ângulo interno de {t['ang_interno']}"
        else:
            frase_deflexao = f"deste ponto, deflete à {t['deflexao']}, formando um ângulo interno de {t['ang_interno']}"

        if t["ponto_atual"] == 1:
            texto_desc += f"do ponto 01, segue com azimute de {t['azimute']} pela distância de {t['distancia']:.2f}m até encontrar o ponto {t['ponto_proximo']:02d}, de coordenadas E= {t['coord_e']}m e N= {t['coord_n']}m, {t['txt_confrontante']}, "
        elif t["ponto_atual"] < st.session_state.num_vertices:
            texto_desc += f"{frase_deflexao}, e segue com azimute de {t['azimute']} pela distância de {t['distancia']:.2f}m até encontrar o ponto {t['ponto_proximo']:02d}, de coordenadas E= {t['coord_e']}m e N= {t['coord_n']}m, {t['txt_confrontante']}, "
        else:
            texto_desc += f"{frase_deflexao}, e segue com azimute de {t['azimute']}, {t['txt_confrontante']} pela distância de {t['distancia']:.2f}m até encontrar o ponto 01, ângulo interno de {ang_fechamento}, origem desta descrição, encerrando com um perímetro de {perimetro}m e uma área total de {area_total} m²."
            
    p2.add_run(texto_desc)
    
    doc.add_paragraph()
    doc.add_paragraph(f"ART n° {art}.")
    
    doc.add_paragraph("\n")
    
    # Ajuste nas Assinaturas para não quebrar linha
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    cell_1 = table.cell(0, 0)
    p_ass1 = cell_1.paragraphs[0]
    p_ass1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass1.add_run("______________________________\nEng. Emerson Alves dos Santos\n(Responsável Técnico)\nCREA-SP n° 506.999.155-2")
    
    cell_2 = table.cell(0, 1)
    p_ass2 = cell_2.paragraphs[0]
    p_ass2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass2.add_run(f"______________________________\n{proprietario}\nRG: {rg}\nCPF: {cpf}")

    # Fonte do rodapé ajustada para Pt(8)
    section = doc.sections[0]
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.text = "Rua Dr. Vicente Giacaglini, 528 – CEP 03203-000 - Vila Bela – São Paulo - e-mail: nickhull.eng@gmail.com - Fone: 55 11 99851-1552"
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_foot.runs:
        run.font.size = Pt(8)
        run.font.name = 'Times New Roman'

    return doc

st.divider()
if st.button("📄 GERAR MEMORIAL DESCRITIVO (WORD)", type="primary", use_container_width=True):
    doc = gerar_documento()
    bio = io.BytesIO()
    doc.save(bio)
    
    st.download_button(
        label="💾 Baixar Documento Oficial",
        data=bio.getvalue(),
        file_name="Memorial_Descritivo_Oficial.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )
    st.success("✅ Memorial gerado! Padrão ABNT e Cartorial aplicados com sucesso.")
